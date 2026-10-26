#!/usr/bin/env python3
"""Genera dashboard Excel de KPI Confiabilidad — versión formateada."""

from openpyxl import Workbook
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, LineChart
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.marker import Marker
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from datetime import date
import json, os, sys

# ── COLORES ────────────────────────────────────────────────────────────────────
C_DARK    = "0F172A"
C_ACCENT  = "1D4ED8"
C_GREEN   = "16A34A"
C_YELLOW  = "D97706"
C_RED     = "DC2626"
C_CYAN    = "0891B2"
C_ORANGE  = "EA580C"
C_WHITE   = "FFFFFF"
C_BG      = "F1F5F9"
C_SURFACE = "FFFFFF"
C_BORDER  = "E2E8F0"
C_TEXT2   = "64748B"
C_HEADER  = "1E293B"
C_GREEN_L = "DCFCE7"
C_YELLOW_L= "FEF9C3"
C_RED_L   = "FEE2E2"
C_BLUE_L  = "DBEAFE"

# ── HELPERS ESTILO ─────────────────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, size=11, color=C_HEADER, name="Calibri", italic=False):
    return Font(bold=bold, size=size, color=color, name=name, italic=italic)

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border(sides="all"):
    s = Side(style="thin", color=C_BORDER)
    n = Side(style=None)
    if sides == "all":
        return Border(left=s, right=s, top=s, bottom=s)
    if sides == "bottom":
        return Border(bottom=s)
    if sides == "top":
        return Border(top=Side(style="medium", color=C_ACCENT))
    return Border()

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header_row(ws, row, values, bg=C_HEADER, fg=C_WHITE, size=10):
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill(bg)
        c.font = font(bold=True, size=size, color=fg)
        c.alignment = align("center")
        c.border = thin_border()

def data_row(ws, row, values, bg=None, bold=False, center_cols=None):
    center_cols = center_cols or []
    bg_c = bg or (C_BG if row % 2 == 0 else C_WHITE)
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill(bg_c)
        c.font = font(bold=bold, size=10)
        c.alignment = align("center" if col in center_cols else "left")
        c.border = thin_border()

def merge_title(ws, row, col_start, col_end, text, bg=C_DARK, fg=C_WHITE, size=13):
    ws.merge_cells(start_row=row, start_column=col_start,
                   end_row=row, end_column=col_end)
    c = ws.cell(row=row, column=col_start, value=text)
    c.fill = fill(bg)
    c.font = font(bold=True, size=size, color=fg)
    c.alignment = align("center")

def nota_fill(nota):
    if nota == 5: return C_GREEN_L, C_GREEN
    if nota == 3: return C_YELLOW_L, C_YELLOW
    if nota == 1: return C_RED_L, C_RED
    return C_BG, C_TEXT2

def disp_fill(disp):
    if disp >= 80: return C_GREEN_L, C_GREEN
    if disp >= 60: return C_YELLOW_L, C_YELLOW
    return C_RED_L, C_RED

# ── DATOS ──────────────────────────────────────────────────────────────────────
KPI_DEF = [
    {"id":1,"nombre":"Cumplimiento PM","freq":"Semanal","meta":90,"pond":20,"inv":False},
    {"id":2,"nombre":"Dotacional","freq":"Mensual","meta":95,"pond":10,"inv":False},
    {"id":3,"nombre":"Traslados","freq":"Semanal","meta":95,"pond":20,"inv":False},
    {"id":4,"nombre":"Rotación Personal","freq":"Mensual","meta":15,"pond":10,"inv":True},
    {"id":5,"nombre":"Disponibilidad Flota","freq":"Mensual","meta":80,"pond":20,"inv":False},
    {"id":6,"nombre":"SOMA","freq":"Mensual","meta":85,"pond":20,"inv":False},
]

DET_SEMANAS = [
    {"id":"S40","label":"Sem 40","eqac":8.4,"pl":26.1,"np":7.7,"rep":3.2,"parque":5},
    {"id":"S41","label":"Sem 41","eqac":5.8,"pl":23.4,"np":7.9,"rep":16.1,"parque":4},
    {"id":"S42","label":"Sem 42","eqac":1.2,"pl":38.3,"np":10.0,"rep":7.8,"parque":4},
    {"id":"S43","label":"Sem 43","eqac":3.2,"pl":12.0,"np":12.5,"rep":12.1,"parque":6},
]

DET_EQUIPOS = {
    "S40":[{"id":"JR001","disp":82,"f":3,"mttr":28},{"id":"JR002","disp":78,"f":4,"mttr":22},
           {"id":"JR003","disp":55,"f":2,"mttr":65},{"id":"JR004","disp":75,"f":3,"mttr":20},
           {"id":"JR005","disp":72,"f":5,"mttr":18},{"id":"JR006","disp":48,"f":4,"mttr":52},
           {"id":"JR007","disp":70,"f":2,"mttr":30},{"id":"JR008","disp":80,"f":2,"mttr":15},
           {"id":"JR009","disp":66,"f":6,"mttr":12},{"id":"JR010","disp":88,"f":1,"mttr":10}],
    "S41":[{"id":"JR001","disp":85,"f":2,"mttr":20},{"id":"JR002","disp":62,"f":6,"mttr":28},
           {"id":"JR003","disp":30,"f":2,"mttr":100},{"id":"JR004","disp":68,"f":4,"mttr":30},
           {"id":"JR005","disp":60,"f":5,"mttr":35},{"id":"JR006","disp":22,"f":3,"mttr":88},
           {"id":"JR007","disp":74,"f":2,"mttr":22},{"id":"JR008","disp":76,"f":3,"mttr":18},
           {"id":"JR009","disp":64,"f":5,"mttr":15},{"id":"JR010","disp":82,"f":2,"mttr":12}],
    "S42":[{"id":"JR001","disp":80,"f":3,"mttr":22},{"id":"JR002","disp":58,"f":7,"mttr":26},
           {"id":"JR003","disp":15,"f":2,"mttr":130},{"id":"JR004","disp":65,"f":4,"mttr":32},
           {"id":"JR005","disp":55,"f":6,"mttr":40},{"id":"JR006","disp":18,"f":3,"mttr":92},
           {"id":"JR007","disp":70,"f":3,"mttr":24},{"id":"JR008","disp":72,"f":3,"mttr":20},
           {"id":"JR009","disp":62,"f":5,"mttr":18},{"id":"JR010","disp":80,"f":2,"mttr":14}],
    "S43":[{"id":"JR001","disp":87,"f":2,"mttr":8},{"id":"JR002","disp":72,"f":5,"mttr":22},
           {"id":"JR003","disp":0,"f":1,"mttr":168},{"id":"JR004","disp":70,"f":3,"mttr":28},
           {"id":"JR005","disp":64,"f":4,"mttr":36},{"id":"JR006","disp":7,"f":3,"mttr":90},
           {"id":"JR007","disp":73,"f":2,"mttr":25},{"id":"JR008","disp":73,"f":3,"mttr":15},
           {"id":"JR009","disp":69,"f":4,"mttr":18},{"id":"JR010","disp":85,"f":1,"mttr":12}],
}

DET_NP = [
    {"s":"Eléctrico","indisp":4.5,"f":3,"hrs":81,
     "comps":"Sensor inclinación JR004/JR009/JR002 | Cable alimentación EX | Profundímetro | Sensor R243/S13"},
    {"s":"Perforación","indisp":4.1,"f":3,"hrs":64,
     "comps":"Cil. avance perforadora JR009 | Mordazas | Acumulador"},
    {"s":"Hidráulico","indisp":2.6,"f":13,"hrs":47,
     "comps":"Flexibles JR002/JR005/JR008"},
    {"s":"Aire Acond.","indisp":1.3,"f":2,"hrs":21,
     "comps":"Compresor | Ductos"},
]

# ── INTENTA LEER REGISTROS EXPORTADOS (JSON opcional) ─────────────────────────
REGISTROS = []
json_path = os.path.join(os.path.dirname(__file__), "kpi_registros.json")
if os.path.exists(json_path):
    with open(json_path, encoding="utf-8") as f:
        REGISTROS = json.load(f)

# ── WORKBOOK ───────────────────────────────────────────────────────────────────
wb = Workbook()
wb.remove(wb.active)  # quita Sheet vacío

# ══════════════════════════════════════════════════════════════════════════════
# HOJA 1: DASHBOARD PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════════
ws = wb.create_sheet("📊 Dashboard")
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = C_ACCENT

set_col_widths(ws, [2, 22, 14, 12, 12, 12, 14, 12, 2])

# Banner título
ws.row_dimensions[1].height = 8
ws.row_dimensions[2].height = 36
ws.row_dimensions[3].height = 20
ws.row_dimensions[4].height = 8

merge_title(ws, 2, 2, 8, "🏗  KPI CONFIABILIDAD CONTRACTUAL", C_DARK, C_WHITE, 16)
ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=8)
c = ws.cell(row=3, column=2, value="Contrato CMC · Codelco · " + date.today().strftime("%d/%m/%Y"))
c.fill = fill(C_ACCENT)
c.font = font(size=10, color=C_WHITE)
c.alignment = align("center")

# ── KPI CARDS (fila 5–12) ──
row = 5
ws.row_dimensions[row].height = 10
row = 6
ws.row_dimensions[row].height = 14

# Header KPIs
merge_title(ws, row, 2, 8, "INDICADORES CONTRACTUALES", C_HEADER, C_WHITE, 10)
row += 1

hdr = ["ID","Indicador","Frecuencia","Meta","Ult. Valor %","Nota (1/3/5)","Contribución","Estado"]
header_row(ws, row, hdr, bg=C_ACCENT)
ws.row_dimensions[row].height = 20
row += 1

# Si tenemos registros reales, calcular último por KPI
def last_by_kpi(registros):
    mapa = {}
    for r in sorted(registros, key=lambda x: x.get("creadoEn","")):
        mapa[r["kpiId"]] = r
    return mapa

mapa_last = last_by_kpi(REGISTROS)
score_sum = 0.0
score_pond = 0

for k in KPI_DEF:
    r = mapa_last.get(k["id"])
    nota = r["nota"] if r else None
    pct  = round(r["pct"], 1) if r else None
    contrib = round(nota * k["pond"] / 100, 2) if nota else None
    estado = ("✅ CUMPLE" if nota == 5 else "⚠️ OBSERVADO" if nota == 3
              else "❌ NO CUMPLE" if nota == 1 else "— Sin datos")

    bg_l, fg_c = nota_fill(nota)
    vals = [k["id"], k["nombre"], k["freq"],
            ("≤" if k["inv"] else "≥") + str(k["meta"]) + "%",
            str(pct) + "%" if pct is not None else "—",
            nota if nota else "—",
            "+" + str(contrib) if contrib else "—",
            estado]

    for col, val in enumerate(vals, 2):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill(bg_l)
        c.font = font(bold=(col == 7), size=10,
                      color=fg_c if col in (7, 9) else C_HEADER)
        c.alignment = align("center" if col in (2, 4, 5, 6, 7, 8, 9) else "left")
        c.border = thin_border()

    if nota:
        score_sum += nota * k["pond"] / 100
        score_pond += k["pond"]

    ws.row_dimensions[row].height = 18
    row += 1

# ── SCORE BANNER ──
row += 1
score_val = round(score_sum, 2) if score_pond else 0.0
score_label = ("EXCELENTE" if score_val >= 4.5 else "BUENO" if score_val >= 3.5
               else "REGULAR" if score_val >= 2.5 else "CRÍTICO" if score_val > 0 else "SIN DATOS")
score_color = (C_GREEN if score_val >= 4 else C_YELLOW if score_val >= 2.5
               else C_RED if score_val > 0 else C_TEXT2)

ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
c = ws.cell(row=row, column=2, value=f"SCORE PONDERADO TOTAL: {score_val} / 5.00")
c.fill = fill(score_color)
c.font = font(bold=True, size=14, color=C_WHITE)
c.alignment = align("center")
ws.row_dimensions[row].height = 28

ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=8)
c = ws.cell(row=row, column=5, value=score_label)
c.fill = fill(score_color)
c.font = font(bold=True, size=14, color=C_WHITE)
c.alignment = align("center")

# ══════════════════════════════════════════════════════════════════════════════
# HOJA 2: SCORE POR PERIODO
# ══════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("📈 Score por Periodo")
ws2.sheet_view.showGridLines = False
ws2.sheet_properties.tabColor = C_GREEN

set_col_widths(ws2, [2, 14] + [14]*len(KPI_DEF) + [14, 12, 2])

merge_title(ws2, 2, 2, 3 + len(KPI_DEF), "EVOLUCIÓN SCORE POR PERIODO", C_DARK, C_WHITE, 14)
ws2.row_dimensions[2].height = 30

row2 = 4
hdr2 = ["Periodo"] + [k["nombre"] for k in KPI_DEF] + ["SCORE TOTAL", "Estado"]
header_row(ws2, row2, hdr2, bg=C_ACCENT)
ws2.row_dimensions[row2].height = 20
row2 += 1

periodos = sorted(set(r["periodo"] for r in REGISTROS)) if REGISTROS else []
chart_data_rows = []

for p in periodos:
    mapa_p = {}
    for r in REGISTROS:
        if r["periodo"] == p:
            mapa_p[r["kpiId"]] = r["nota"]
    sc = sum(mapa_p.get(k["id"], 0) * k["pond"] / 100 for k in KPI_DEF
             if mapa_p.get(k["id"]))
    sc = round(sc, 2)
    estado2 = ("EXCELENTE" if sc >= 4.5 else "BUENO" if sc >= 3.5
               else "REGULAR" if sc >= 2.5 else "CRÍTICO" if sc > 0 else "—")
    bg_s, _ = nota_fill(5 if sc >= 4 else 3 if sc >= 2.5 else 1 if sc > 0 else None)

    row_vals = [p] + [mapa_p.get(k["id"], "") for k in KPI_DEF] + [sc, estado2]
    for col, val in enumerate(row_vals, 2):
        c = ws2.cell(row=row2, column=col, value=val)
        nota_v = val if col > 2 and col <= 2 + len(KPI_DEF) and isinstance(val, int) else None
        bg_c2, fg_c2 = nota_fill(nota_v)
        c.fill = fill(bg_c2 if nota_v else (C_BG if row2 % 2 == 0 else C_WHITE))
        c.font = font(bold=(col == 2 + len(KPI_DEF) + 1), size=10,
                      color=fg_c2 if nota_v else C_HEADER)
        c.alignment = align("center")
        c.border = thin_border()
    chart_data_rows.append(sc)
    ws2.row_dimensions[row2].height = 18
    row2 += 1

# Si no hay registros reales, fila indicativa
if not periodos:
    ws2.cell(row=row2, column=2, value="— Sin registros. Ingresa datos en el dashboard. —")
    ws2.cell(row=row2, column=2).font = font(italic=True, color=C_TEXT2)

# ══════════════════════════════════════════════════════════════════════════════
# HOJA 3: DETENCIONES POR SEMANA
# ══════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("🔧 Detenciones")
ws3.sheet_view.showGridLines = False
ws3.sheet_properties.tabColor = C_RED

set_col_widths(ws3, [2, 12, 12, 13, 13, 16, 14, 14, 16, 2])

merge_title(ws3, 2, 2, 9, "DETENCIONES JUMBOS RADIALES — RESUMEN SEMANAL", C_DARK, C_WHITE, 14)
ws3.row_dimensions[2].height = 30

row3 = 4
header_row(ws3, row3, ["Semana","EQAC %","Mant. PL %","Mant. NP %",
                        "Falta Repuesto %","Total InDisp %","Disponibilidad %",
                        "Parque Operativo"], bg=C_RED)
ws3.row_dimensions[row3].height = 20
row3 += 1

for s in DET_SEMANAS:
    total = round(s["eqac"] + s["pl"] + s["np"] + s["rep"], 1)
    disp  = round(100 - total, 1)
    bg_d, fg_d = disp_fill(disp)
    vals3 = [s["label"], s["eqac"], s["pl"], s["np"], s["rep"], total, disp, s["parque"]]
    for col, val in enumerate(vals3, 2):
        c = ws3.cell(row=row3, column=col, value=val)
        if col == 8:  # disponibilidad
            c.fill = fill(bg_d)
            c.font = font(bold=True, size=10, color=fg_d)
        elif col == 7:  # total indisp
            c.fill = fill(C_RED_L)
            c.font = font(bold=True, size=10, color=C_RED)
        else:
            c.fill = fill(C_BG if row3 % 2 == 0 else C_WHITE)
            c.font = font(size=10)
        c.alignment = align("center")
        c.border = thin_border()
    ws3.row_dimensions[row3].height = 18
    row3 += 1

# Fila totales
row3 += 1
totals3 = ["PROMEDIO",
           round(sum(s["eqac"] for s in DET_SEMANAS)/len(DET_SEMANAS),1),
           round(sum(s["pl"]   for s in DET_SEMANAS)/len(DET_SEMANAS),1),
           round(sum(s["np"]   for s in DET_SEMANAS)/len(DET_SEMANAS),1),
           round(sum(s["rep"]  for s in DET_SEMANAS)/len(DET_SEMANAS),1),
           None, None, None]
for col, val in enumerate(totals3, 2):
    if val is None: continue
    c = ws3.cell(row=row3, column=col, value=val)
    c.fill = fill(C_HEADER)
    c.font = font(bold=True, size=10, color=C_WHITE)
    c.alignment = align("center")
    c.border = thin_border()

# ── Sub-tabla NP por Sistema ──
row3 += 3
merge_title(ws3, row3, 2, 9, "DETALLE NP POR SISTEMA", C_HEADER, C_WHITE, 11)
row3 += 1
header_row(ws3, row3, ["Sistema","% InDisp NP","N° Fallas","Hrs Detención",
                        "Componentes Afectados","","",""], bg=C_HEADER)
ws3.row_dimensions[row3].height = 18
row3 += 1

NP_BG = {"Eléctrico": C_BLUE_L, "Perforación": C_GREEN_L,
          "Hidráulico": C_YELLOW_L, "Aire Acond.": "FCE7F3"}

for n in DET_NP:
    bg_n = NP_BG.get(n["s"], C_BG)
    vals_n = [n["s"], n["indisp"], n["f"], n["hrs"], n["comps"], "", "", ""]
    for col, val in enumerate(vals_n, 2):
        c = ws3.cell(row=row3, column=col, value=val)
        c.fill = fill(bg_n)
        c.font = font(bold=(col == 2), size=10,
                      color=C_RED if col == 3 else C_HEADER)
        c.alignment = align("left" if col == 6 else "center", wrap=(col == 6))
        c.border = thin_border()
    ws3.row_dimensions[row3].height = 30
    row3 += 1

# Fila total NP
tot_np = {"indisp": round(sum(n["indisp"] for n in DET_NP), 1),
          "f": sum(n["f"] for n in DET_NP),
          "hrs": sum(n["hrs"] for n in DET_NP)}
for col, val in enumerate(["TOTAL NP", tot_np["indisp"], tot_np["f"], tot_np["hrs"], "", "", "", ""], 2):
    c = ws3.cell(row=row3, column=col, value=val)
    c.fill = fill(C_DARK)
    c.font = font(bold=True, size=10, color=C_WHITE)
    c.alignment = align("center")
    c.border = thin_border()

# ══════════════════════════════════════════════════════════════════════════════
# HOJA 4: EQUIPOS (JACK KNIFE DATA)
# ══════════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("⚙️ Equipos")
ws4.sheet_view.showGridLines = False
ws4.sheet_properties.tabColor = C_CYAN

set_col_widths(ws4, [2, 10, 10, 16, 12, 14, 12, 2])

merge_title(ws4, 2, 2, 7, "DISPONIBILIDAD Y FALLAS POR EQUIPO", C_DARK, C_WHITE, 14)
ws4.row_dimensions[2].height = 30

row4 = 4
header_row(ws4, row4, ["Semana","Equipo","Disponibilidad %","N° Fallas",
                        "MTTR (hrs)","Estado"], bg=C_CYAN)
ws4.row_dimensions[row4].height = 20
row4 += 1

prev_sem = None
for s in DET_SEMANAS:
    total_indisp = round(s["eqac"]+s["pl"]+s["np"]+s["rep"],1)
    disp_flota = round(100 - total_indisp, 1)

    # Separador de semana
    ws4.merge_cells(start_row=row4, start_column=2, end_row=row4, end_column=7)
    c = ws4.cell(row=row4, column=2,
                 value=f"{s['label']}  ·  Disp. Flota: {disp_flota}%  ·  Parque: {s['parque']} equipos")
    c.fill = fill(C_HEADER)
    c.font = font(bold=True, size=10, color=C_WHITE)
    c.alignment = align("left")
    ws4.row_dimensions[row4].height = 16
    row4 += 1

    for e in DET_EQUIPOS[s["id"]]:
        bg_e, fg_e = disp_fill(e["disp"])
        estado_e = "✅ OK" if e["disp"] >= 80 else "⚠️ ATENCIÓN" if e["disp"] >= 60 else "❌ CRÍTICO"
        vals4 = [s["label"], e["id"], e["disp"], e["f"], e["mttr"], estado_e]
        for col, val in enumerate(vals4, 2):
            c = ws4.cell(row=row4, column=col, value=val)
            if col == 4:  # disponibilidad
                c.fill = fill(bg_e)
                c.font = font(bold=True, size=10, color=fg_e)
            elif col == 7:  # estado
                c.fill = fill(bg_e)
                c.font = font(bold=False, size=10, color=fg_e)
            else:
                c.fill = fill(C_BG if row4 % 2 == 0 else C_WHITE)
                c.font = font(size=10)
            c.alignment = align("center")
            c.border = thin_border()
        ws4.row_dimensions[row4].height = 16
        row4 += 1
    row4 += 1

# ══════════════════════════════════════════════════════════════════════════════
# HOJA 5: HISTORIAL KPIs (datos reales si existen)
# ══════════════════════════════════════════════════════════════════════════════
ws5 = wb.create_sheet("📋 Historial KPIs")
ws5.sheet_view.showGridLines = False
ws5.sheet_properties.tabColor = "7C3AED"

set_col_widths(ws5, [2, 5, 28, 11, 13, 12, 12, 8, 7, 10, 10, 28, 2])

merge_title(ws5, 2, 2, 12, "HISTORIAL DE REGISTROS KPI", C_DARK, C_WHITE, 14)
ws5.row_dimensions[2].height = 30

row5 = 4
header_row(ws5, row5, ["N°","Indicador","Frecuencia","Periodo",
                        "Realizado","Planificado","%","Nota",
                        "Pond%","Contribución","Observación"], bg="7C3AED")
ws5.row_dimensions[row5].height = 20
row5 += 1

kpi_map = {k["id"]: k for k in KPI_DEF}

if REGISTROS:
    for i, r in enumerate(sorted(REGISTROS, key=lambda x: x.get("creadoEn","")), 1):
        k = kpi_map.get(r.get("kpiId"))
        nota = r.get("nota")
        bg_h, fg_h = nota_fill(nota)
        vals5 = [i, k["nombre"] if k else "—", k["freq"] if k else "—",
                 r.get("periodo",""), r.get("num",""), r.get("den",""),
                 round(r.get("pct",0),1), nota,
                 k["pond"] if k else "", round(nota*k["pond"]/100,3) if nota and k else "",
                 r.get("obs","")]
        for col, val in enumerate(vals5, 2):
            c = ws5.cell(row=row5, column=col, value=val)
            c.fill = fill(bg_h if col == 9 else (C_BG if row5%2==0 else C_WHITE))
            c.font = font(bold=(col==9), size=10,
                          color=fg_h if col==9 else C_HEADER)
            c.alignment = align("center" if col in (2,4,5,6,8,9,10,11) else "left")
            c.border = thin_border()
        ws5.row_dimensions[row5].height = 16
        row5 += 1
else:
    ws5.merge_cells(start_row=row5, start_column=2, end_row=row5, end_column=12)
    c = ws5.cell(row=row5, column=2,
                 value="Sin registros. Exporta kpi_registros.json desde el dashboard para incluir datos reales.")
    c.font = font(italic=True, color=C_TEXT2, size=11)
    c.alignment = align("center")

# ── GUARDAR ────────────────────────────────────────────────────────────────────
output = os.path.join(os.path.dirname(__file__),
                      f"Dashboard_KPI_Confiabilidad_{date.today().strftime('%Y-%m-%d')}.xlsx")
wb.save(output)
print(f"✅ Excel generado: {output}")
